"""Tiny synthetic workbooks for Phase 6 tests. Never use client files."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


def write_books_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Books"
    sheet.append(["Book", "Name"])
    sheet.append(["A", "Alpha"])
    sheet.append(["B", "Bravo"])
    workbook.save(path)
    workbook.close()
    return path


def write_rates_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Date", "Book", "Rate", "Notes"])
    sheet.append(["2026-01-01", "A", 1.5, None])
    sheet.append(["2026-01-02", "B", 2.0, None])
    sheet["E2"] = "=[Books.xlsx]Books!B2"
    sheet["E3"] = "=C3*1"
    workbook.defined_names.add(
        DefinedName(name="RateTable", attr_text="Rates!$A$1:$C$3")
    )
    workbook.save(path)
    workbook.close()
    return path


def write_wacc_pack_xlsx(path: Path) -> Path:
    """Damodaran-shaped single workbook: title row, WACC pointer, empty cols."""
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Input sheet"
    inputs["A1"] = (
        "FCFF Simple Ginzu — illustration workbook for cost of capital tests"
    )
    inputs["A2"] = "Item"
    inputs["B2"] = "Value"
    inputs["C2"] = "Note"
    for row in range(3, 11):
        inputs[f"A{row}"] = f"Input {row}"
        inputs[f"B{row}"] = row
        inputs[f"C{row}"] = "keep"
    inputs["A20"] = "Industry lookup"
    inputs["B20"] = "=VLOOKUP(B3,'Cost of capital'!A10:C13,2,FALSE)"
    inputs["A35"] = "Initial cost of capital"
    inputs["B35"] = "='Cost of capital'!B13"
    inputs["J40"] = ""

    cost = workbook.create_sheet("Cost of capital")
    cost["A10"] = "Equity weight"
    cost["B10"] = 0.6
    cost["A11"] = "Cost of equity"
    cost["B11"] = 0.1
    cost["A12"] = "Debt component"
    cost["B12"] = 0.02
    cost["A13"] = "Cost of capital"
    cost["B13"] = "=B10*B11+B12"

    valuation = workbook.create_sheet("Valuation output")
    valuation["A1"] = "Year"
    valuation["B1"] = "FCFF"
    valuation["A20"] = "Terminal value"
    valuation["F20"] = 100
    valuation["G20"] = 0.03
    valuation["L20"] = "=F20*(1+G20)/('Cost of capital'!B13-G20)"
    valuation["P22"] = ""

    workbook.defined_names.add(
        DefinedName(name="RateTable", attr_text="'Input sheet'!$A$2:$C$10")
    )
    workbook.save(path)
    workbook.close()
    return path


def write_mixed_csv(path: Path) -> Path:
    path.write_text(
        "id,amount,note,spare\n"
        "1,10,hello,\n"
        "2,x,world,\n",
        encoding="utf-8",
    )
    return path


def write_agency_costs_xlsx(
    path: Path,
    *,
    early_count: int = 8,
    late_count: int = 6,
    early_amount: float = 1,
    late_amount: float = 1000,
    max_row_amount: float = 4000,
) -> Path:
    """Same layout as `write_agency_costs_csv`, native xlsx, no formulas."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "agency_costs"
    sheet.append(["agency", "amount", "year"])
    for _ in range(early_count):
        sheet.append(["Alpha", early_amount, 2010])
    for _ in range(late_count):
        sheet.append(["Zulu", late_amount, 2020])
    sheet.append(["Omega", max_row_amount, 2021])
    workbook.save(path)
    workbook.close()
    return path


def write_agency_costs_csv(
    path: Path,
    *,
    early_count: int = 8,
    late_count: int = 6,
    early_amount: float = 1,
    late_amount: float = 1000,
    max_row_amount: float = 4000,
) -> Path:
    """Synthetic agency table. Late names sit after `early_count` rows.

    Default: Alpha sum is small, Zulu has the largest *sum*, Omega has
    the largest *single row*. Used to prove max-row ≠ group total.
    """
    lines = ["agency,amount,year"]
    for _ in range(early_count):
        lines.append(f"Alpha,{_csv_number(early_amount)},2010")
    for _ in range(late_count):
        lines.append(f"Zulu,{_csv_number(late_amount)},2020")
    lines.append(f"Omega,{_csv_number(max_row_amount)},2021")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_agency_and_subagency_csv(path: Path) -> Path:
    """Agency vs sub_agency: 'list the agencies' must not fail closed."""
    path.write_text(
        "agency,sub_agency,amount,year\n"
        "Alpha,Alpha-1,10,2010\n"
        "Alpha,Alpha-2,20,2011\n"
        "Zulu,Zulu-1,1000,2020\n"
        "Zulu,Zulu-2,2000,2020\n"
        "Omega,Omega-1,400,2021\n",
        encoding="utf-8",
    )
    return path


def write_many_agencies_csv(path: Path, count: int = 48) -> Path:
    """One row per agency so all names must appear on the stats card."""
    lines = ["agency,amount,year"]
    for index in range(1, count + 1):
        lines.append(f"Agency_{index:02d},{index},2015")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _csv_number(value: float) -> str:
    if float(value).is_integer():
        return str(int(value))
    return str(value)


def write_tiny_xls(path: Path) -> Path:
    """Synthetic BIFF2 workbook. No client files; no xlwt binary."""
    import struct

    chunks: list[bytes] = []

    def record(code: int, payload: bytes) -> None:
        chunks.append(struct.pack("<HH", code, len(payload)))
        chunks.append(payload)

    def label(row: int, col: int, text: str) -> None:
        raw = text.encode("latin-1")
        record(0x0004, struct.pack("<HHB", row, col, 0) + bytes([len(raw)]) + raw)

    def number(row: int, col: int, value: float) -> None:
        record(0x0003, struct.pack("<HHB", row, col, 0) + struct.pack("<d", float(value)))

    record(0x0009, struct.pack("<HH", 0x0002, 0x0010))
    record(0x0000, struct.pack("<HHBB", 0, 3, 0, 3) + b"\x00")
    label(0, 0, "id")
    label(0, 1, "amount")
    label(0, 2, "note")
    number(1, 0, 1)
    number(1, 1, 10)
    label(1, 2, "hello")
    number(2, 0, 2)
    number(2, 1, 20)
    label(2, 2, "world")
    record(0x000A, b"")
    path.write_bytes(b"".join(chunks))
    return path


def add_zip_members(path: Path, members: dict[str, bytes]) -> Path:
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(tmp, "w") as dest:
        for info in source.infolist():
            dest.writestr(info, source.read(info.filename))
        for name, data in members.items():
            dest.writestr(name, data)
    tmp.replace(path)
    return path
